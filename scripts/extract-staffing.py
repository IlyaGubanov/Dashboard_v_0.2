"""Извлекает только агрегаты выбранного штатного расписания, без персональных данных."""
import json
import re
import sys
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
target = Path(sys.argv[2])
book = openpyxl.load_workbook(source, data_only=True, read_only=True)
sheet = book['АНО МосводостокСтройТрест']
data = list(sheet.iter_rows(min_row=6, max_row=46, max_col=22, values_only=True))
number = lambda v: v if isinstance(v, (int, float)) else None
periods = []
for date, plan_col, actual_col in [('2025-01-01', 6, 14), ('2025-07-01', 7, 15), ('2026-01-01', 9, 16), ('2026-07-01', 10, 17)]:
    rows = [dict(id=f'source-{r[0]}', department=r[1], category=r[2], planned=number(r[plan_col]), actual=number(r[actual_col]), note='', isArchived=False, function=r[21] or '') for r in data]
    periods.append(dict(month=date[:7], asOf=date, note=f'Источник: {source.name}, {sheet.title}, строки 6–46. Пустые ячейки и «нет» сохранены как отсутствие данных.', rows=rows))
# Срез на 26.10.2026 — только итоги по блокам из оперативной сводки; штата по блокам в ней нет, ФИО руководителя из названия убираем.
svod = book['СВОД по подразделениям c 20.10.']
cells = dict(enumerate(svod.iter_rows(min_row=1, max_row=20, max_col=3, values_only=True), start=1))
assert cells[4][2].date().isoformat() == '2026-10-26', cells[4][2]
groups = {6: 'ГУП МВС (СМУ+РСУ+ТЛС)', 7: 'РЭБ (без ТЛС)', 8: 'РСУ СР+УСР от Учредителя', 9: 'РСУ-1', 10: 'РСУ-2', 11: 'РСУ-3', 12: 'РСУ-5', 13: 'РСУ-6', 14: 'АУП',
          15: 'Управление гражданского строительства', 16: 'Управление обеспечения и эксплуатации производственной базы'}
rows = [dict(id=f'source-svod-{r}', department=re.sub(r'\s*\(Рук\.[^)]*\)', '', str(cells[r][0])).strip(), category='АУП' if r == 14 else 'ВП' if r >= 15 else 'ОП', planned=None, actual=cells[r][2], note='', isArchived=False, function='', group=g) for r, g in groups.items()]
assert sum(r['actual'] for r in rows) == cells[17][2], 'сумма блоков не совпала с итогом строки 17'
space = lambda n: f'{n:,}'.replace(',', ' ')
periods.append(dict(month='2026-10', asOf='2026-10-26', note=f'Источник: {source.name}, «{svod.title}», столбец на 26.10.2026 (количество сотрудников по 1С), строки 6–16. Штат по блокам в источнике не указан: утверждено всего {space(cells[18][2])} штатных единиц, вакансии — {space(cells[20][2])}.', rows=rows))
payroll = [dict(year=y, label='С октября 2024' if y == 2024 else f'{y} — период в источнике не уточнён', rows=[dict(id=f'source-{r[0]}', department=r[1], category=r[2], amount=number(r[c])) for r in data]) for y,c in [(2024,18),(2025,19),(2026,20)]]
result = dict(source=source.name, sheet=sheet.title, sheetCount=len(book.sheetnames), periods=periods, payroll=payroll)
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(dict(sheets=len(book.sheetnames), departments=len(data), periods=[dict(month=p['month'], planned=sum(r['planned'] or 0 for r in p['rows']), actual=sum(r['actual'] or 0 for r in p['rows']), missingPlan=sum(r['planned'] is None for r in p['rows']), missingActual=sum(r['actual'] is None for r in p['rows'])) for p in periods], payroll=[dict(year=p['year'], amount=round(sum(r['amount'] or 0 for r in p['rows']),2)) for p in payroll]), ensure_ascii=False))
